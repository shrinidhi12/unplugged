"""Turn an export of the collaborator sheet into the site's venue data.

    npm run venues:sync                       # newest .xlsx in sheet-export/
    npm run venues:sync -- --input FILE.xlsx  # a specific export
    npm run venues:sync -- --no-geocode       # skip address lookups

Steps: validate every row against the Field guide tab, create missing Venue
IDs, geocode rows that have an address but no coordinates, rebuild the Off
Tech history from pipeline/, then write src/data/venues.json with only
Approved rows and only the columns the Field guide marks "Shown on site = Yes".
Contact columns are dropped unless the row says "OK to show contact publicly?"
= Yes. Nothing is ever written back to the sheet.
"""
from __future__ import annotations

import argparse, datetime, json, math, re, sys, time, unicodedata, urllib.parse, urllib.request
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(HERE / "pipeline"))

EXPORT_DIR = HERE / "sheet-export"
SEED = HERE / "data" / "venues_master_seed.xlsx"
GEOCODE_CACHE = HERE / "data" / "geocode_cache.json"
OUT = ROOT / "src" / "data" / "venues.json"

NOMINATIM = "https://nominatim.openstreetmap.org/search"
USER_AGENT = "UnpluggMe-venue-sync/1.0 (+https://github.com/shrinidhi12/unplugged)"
NYC_BOUNDS = (40.47, 40.93, -74.27, -73.68)  # lat min/max, lng min/max
MOVE_WARN_M = 150

# Sheet column -> key in venues.json. Every "Shown on site = Yes" column must
# be listed; sync fails if the Field guide adds one this file doesn't know.
PUBLIC_KEYS = {
    "Venue name": "name",
    "Venue type": "type",
    "Street address": "address",
    "Neighborhood": "neighborhood",
    "Borough": "borough",
    "Latitude": "lat",
    "Longitude": "lng",
    "Location precision": "precision",
    "Indoor / outdoor": "indoorOutdoor",
    "Capacity, standing (est.)": "capacityStanding",
    "Capacity, seated (est.)": "capacitySeated",
    "Seating": "seating",
    "Tables available": "tables",
    "Private room / buyout": "privateRoom",
    "Weather backup": "weatherBackup",
    "Noise level": "noise",
    "Amplified sound allowed": "amplifiedSound",
    "Phone-free friendly": "phoneFree",
    "Cost": "cost",
    "Cost notes": "costNotes",
    "How to book": "howToBook",
    "Permit notes": "permitNotes",
    "Booking lead time": "leadTime",
    "Hours / best times": "bestTimes",
    "Food & drink": "food",
    "Alcohol": "alcohol",
    "All ages": "allAges",
    "Restrooms": "restrooms",
    "Step-free access": "stepFree",
    "Accessible restroom": "accessibleRestroom",
    "Accessibility notes": "accessibilityNotes",
    "Power outlets": "power",
    "Sound system / mic": "soundSystem",
    "Projector / screen": "projector",
    "Nearest transit": "transit",
    "Website": "website",
    "Instagram": "instagram",
    "Public booking / inquiry link": "bookingLink",
    "Notes for hosts": "notes",
    "Last verified": "lastVerified",
}
# Published only for rows where the venue agreed (OK to show contact = Yes).
# "Relationship notes" is never published.
CONTACT_KEYS = {
    "Key contact": "name",
    "Contact role": "role",
    "Contact email": "email",
    "Contact phone": "phone",
}
OK_CONTACT = "OK to show contact publicly?"
# Columns the site filters on; their allowed values come from the Field guide.
FILTER_COLS = ["Venue type", "Borough", "Indoor / outdoor", "Cost"]
REQUIRED = ["Venue name", "Venue type", "Borough", "Status"]
INTEGER_COLS = ["Capacity, standing (est.)", "Capacity, seated (est.)"]
URL_COLS = ["Website", "Public booking / inquiry link"]
# Venue IDs that would collide with pages under /find-a-venue/.
RESERVED_IDS = {"suggest"}


class Problems:
    """Collects errors and warnings as 'Row N, column: message'."""

    def __init__(self):
        self.errors, self.warnings = [], []

    def add(self, row, col, msg, blocking=True):
        where = f"Row {row}" + (f", {col}" if col else "")
        (self.errors if blocking else self.warnings).append(f"{where}: {msg}")


# ---- Reading the sheet -------------------------------------------------------

def clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def read_field_guide(wb):
    ws = wb["Field guide"]
    it = ws.iter_rows(values_only=True)
    header = [clean(c) for c in next(it)]
    guide = {}
    for r in it:
        d = dict(zip(header, (clean(c) for c in r)))
        if not d.get("Column"):
            continue
        allowed = d.get("Allowed values")
        guide[d["Column"]] = {
            "section": d.get("Section"),
            "shown": d.get("Shown on site?"),
            "allowed": [a.strip() for a in re.split(r"\s*·\s*", allowed)] if allowed else None,
        }
    return guide


def read_venues(wb, guide, problems):
    ws = wb["Venues"]
    it = ws.iter_rows(values_only=True)
    header = [clean(c) for c in next(it)]
    missing = [c for c in guide if c not in header]
    extra = [c for c in header if c and c not in guide]
    if missing:
        problems.add(1, None, f"Venues tab is missing column(s) from the Field guide: {', '.join(missing)}")
    if extra:
        problems.add(1, None, f"Venues tab has column(s) the Field guide doesn't describe: {', '.join(extra)}")
    rows = []
    for i, r in enumerate(it, start=2):
        d = {h: clean(v) for h, v in zip(header, r) if h}
        if any(v is not None for v in d.values()):
            d["_row"] = i
            rows.append(d)
    return rows


# ---- Validation --------------------------------------------------------------

def slugify(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"\(.*?\)", "", s.lower()).replace("'", "")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def as_int(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v) if float(v).is_integer() else None
    m = re.fullmatch(r"~?\s*(\d+)\s*\+?", str(v))
    return int(m.group(1)) if m else None


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    try:
        return datetime.date.fromisoformat(str(v))
    except ValueError:
        return None


def validate(rows, guide, problems):
    seen_ids = {}
    for d in rows:
        n = d["_row"]
        approved = d.get("Status") == "Approved"
        # Problems on rows that won't be published are warnings, not blockers.
        block = approved or d.get("Status") is None

        for col in REQUIRED:
            if d.get(col) is None:
                problems.add(n, col, "required, but blank", block)

        for col, g in guide.items():
            v = d.get(col)
            if v is None or not g["allowed"]:
                continue
            match = next((a for a in g["allowed"] if a.lower() == str(v).lower()), None)
            if match is None:
                problems.add(n, col, f"'{v}' isn't an allowed value ({' · '.join(g['allowed'])})", block)
            else:
                d[col] = match

        vid = d.get("Venue ID")
        if vid is not None:
            if not re.fullmatch(r"[a-z0-9]+(-[a-z0-9]+)*", str(vid)):
                problems.add(n, "Venue ID", f"'{vid}' should be lowercase letters, numbers and dashes", block)
            elif vid in RESERVED_IDS:
                problems.add(n, "Venue ID", f"'{vid}' is reserved for a site page; pick another", block)
            elif vid in seen_ids:
                problems.add(n, "Venue ID", f"'{vid}' is also used on row {seen_ids[vid]}")
            else:
                seen_ids[vid] = n

        lat, lng = d.get("Latitude"), d.get("Longitude")
        if (lat is None) != (lng is None):
            problems.add(n, "Latitude/Longitude", "fill in both or neither", block)
        elif lat is not None:
            try:
                lat, lng = float(lat), float(lng)
            except (TypeError, ValueError):
                problems.add(n, "Latitude/Longitude", f"'{d['Latitude']}, {d['Longitude']}' aren't numbers", block)
            else:
                d["Latitude"], d["Longitude"] = lat, lng
                if not in_nyc(lat, lng):
                    hint = " (swapped?)" if in_nyc(lng, lat) else ""
                    problems.add(n, "Latitude/Longitude", f"{lat}, {lng} isn't in New York City{hint}", block)
        elif d.get("Street address") is None:
            problems.add(n, "Street address", "needs a street address or coordinates", block)

        for col in INTEGER_COLS:
            v = d.get(col)
            if v is None:
                continue
            i = as_int(v)
            if i is None or i <= 0 or i > 50000:
                problems.add(n, col, f"'{v}' should be a whole number of people, e.g. 40", block)
            else:
                d[col] = i

        for col in URL_COLS:
            v = d.get(col)
            if v is not None and not re.match(r"https?://\S+$", str(v)):
                problems.add(n, col, f"'{v}' should be a full link starting with https://", block)
        ig = d.get("Instagram")
        if ig is not None and not re.match(r"(@?[A-Za-z0-9._]+|https?://(www\.)?instagram\.com/\S+)$", str(ig)):
            problems.add(n, "Instagram", f"'{ig}' should be @handle or an instagram.com link", block)

        lv = d.get("Last verified")
        if lv is not None:
            dt = as_date(lv)
            if dt is None:
                problems.add(n, "Last verified", f"'{lv}' should be a date like 2026-10-28", block)
            else:
                d["Last verified"] = dt.isoformat()

    return seen_ids


def in_nyc(lat, lng):
    a, b, c, e = NYC_BOUNDS
    return a <= lat <= b and c <= lng <= e


def assign_ids(rows, seen_ids):
    """Give rows without a Venue ID a unique slug. Returns [(row, name, id)]."""
    created = []
    for d in rows:
        if d.get("Venue ID") is None and d.get("Venue name"):
            base = slugify(str(d["Venue name"])) or "venue"
            vid, k = base, 2
            while vid in seen_ids or vid in RESERVED_IDS:
                vid, k = f"{base}-{k}", k + 1
            seen_ids[vid] = d["_row"]
            d["Venue ID"] = vid
            created.append((d["_row"], d["Venue name"], vid))
    return created


# ---- Geocoding ---------------------------------------------------------------

def metres(a, b):
    (la1, lo1), (la2, lo2) = a, b
    x = math.radians(lo2 - lo1) * math.cos(math.radians((la1 + la2) / 2))
    y = math.radians(la2 - la1)
    return 6371000 * math.hypot(x, y)


def geocode(rows, problems, enabled):
    """Fill coordinates from the address via Nominatim (1 request/sec), cached."""
    cache = json.loads(GEOCODE_CACHE.read_text("utf-8")) if GEOCODE_CACHE.exists() else {}
    looked_up, last_call = [], 0.0
    for d in rows:
        addr = d.get("Street address")
        if d.get("Latitude") is not None or addr is None:
            continue
        q = addr if re.search(r"new york|ny\b|\d{5}", addr, re.I) else f"{addr}, {d.get('Borough') or ''}, New York, NY"
        key = q.lower()
        if key not in cache:
            if not enabled:
                problems.add(d["_row"], "Latitude/Longitude", "blank, and geocoding is off (--no-geocode)",
                             d.get("Status") == "Approved")
                continue
            time.sleep(max(0, 1.1 - (time.time() - last_call)))
            url = NOMINATIM + "?" + urllib.parse.urlencode(
                {"q": q, "format": "jsonv2", "limit": 1, "countrycodes": "us",
                 "viewbox": "-74.27,40.93,-73.68,40.47", "bounded": 1})
            req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            try:
                with urllib.request.urlopen(req, timeout=20) as r:
                    hits = json.load(r)
            except OSError as e:
                problems.add(d["_row"], "Street address", f"geocoding failed ({e}); try again later")
                continue
            finally:
                last_call = time.time()
            cache[key] = [round(float(hits[0]["lat"]), 6), round(float(hits[0]["lon"]), 6)] if hits else None
        hit = cache[key]
        if hit is None:
            problems.add(d["_row"], "Street address",
                         f"couldn't find '{addr}' on OpenStreetMap; check it or enter coordinates",
                         d.get("Status") == "Approved")
            continue
        d["Latitude"], d["Longitude"] = hit
        d.setdefault("Location precision", "exact")
        looked_up.append((d["_row"], d.get("Venue ID"), addr, hit))
    GEOCODE_CACHE.write_text(json.dumps(cache, indent=1, ensure_ascii=False, sort_keys=True) + "\n", "utf-8")
    return looked_up


# ---- History -----------------------------------------------------------------

def load_history():
    import history
    out, candidates = history.build()
    history.write(out)

    def counts(s, sep):
        # "Reading, poetry & book clubs (12), Music & dance (9)": names can hold
        # commas, so match "name (n)" pairs rather than splitting on the separator.
        return [{"activity": m.group(1), "count": int(m.group(2))}
                for m in re.finditer(rf"(.+?) \((\d+)\)(?:{re.escape(sep)}|$)", s or "")]

    by_id = {}
    for r in out:
        host = r["Most common host"]
        by_id[r["Venue ID"]] = {
            "listings": r["Off Tech listings"],
            "issues": r["Issues mentioning"],
            "firstIssue": r["First issue"],
            "latestIssue": r["Latest issue"],
            "latestEvent": r["Latest event"],
            "topActivities": counts(r["Top activities"], ", "),
            "activities": counts(r["All activities"], "; "),
            "signatureEvents": [s.strip() for s in r["Signature events"].split(", ") if s.strip()],
            "topHost": None if host.startswith("—") else host,
        }
    return by_id, candidates


# ---- Publishing --------------------------------------------------------------

def publish(rows, guide, history):
    shown = {c for c, g in guide.items() if g["shown"] == "Yes"}
    unknown = shown - set(PUBLIC_KEYS)
    if unknown:
        sys.exit(f"The Field guide marks {', '.join(sorted(unknown))} as shown on the site, "
                 "but sync.py doesn't know where to put it. Add it to PUBLIC_KEYS.")
    venues = []
    for d in rows:
        if d.get("Status") != "Approved":
            continue
        v = {"id": d["Venue ID"]}
        for col, key in PUBLIC_KEYS.items():
            # A column the Field guide no longer marks as shown stays out.
            v[key] = d.get(col) if col in shown else None
            if isinstance(v[key], str) and key not in ("notes", "accessibilityNotes", "permitNotes", "costNotes"):
                v[key] = re.sub(r"\s+", " ", v[key])
        if v["precision"] is None:
            v["precision"] = "exact"
        if d.get(OK_CONTACT) == "Yes":
            contact = {k: d.get(c) for c, k in CONTACT_KEYS.items() if d.get(c) is not None}
            if contact:
                v["contact"] = contact
        v["history"] = history.get(v["id"])
        venues.append(v)
    venues.sort(key=lambda v: (-(v["history"] or {}).get("listings", 0), v["name"].lower()))
    return venues


def summarize_changes(old, new):
    old_by, new_by = {v["id"]: v for v in old}, {v["id"]: v for v in new}
    lines = []
    for vid in sorted(new_by.keys() - old_by.keys()):
        lines.append(f"  + added    {vid}")
    for vid in sorted(old_by.keys() - new_by.keys()):
        lines.append(f"  - removed  {vid}")
    moved = []
    for vid in sorted(new_by.keys() & old_by.keys()):
        a, b = old_by[vid], new_by[vid]
        changed = [k for k in sorted(set(a) | set(b)) if a.get(k) != b.get(k)]
        if changed:
            lines.append(f"  ~ changed  {vid}: {', '.join(changed)}")
        if None not in (a.get("lat"), b.get("lat")):
            m = metres((a["lat"], a["lng"]), (b["lat"], b["lng"]))
            if m > MOVE_WARN_M:
                moved.append(f"  ! {vid} moved {m:,.0f} m. Check the pin before publishing.")
    return lines, moved


def pick_input(arg):
    if arg:
        return Path(arg)
    exports = sorted(EXPORT_DIR.glob("*.xlsx"), key=lambda p: p.stat().st_mtime)
    exports = [p for p in exports if not p.name.startswith("~$")]
    if not exports:
        sys.exit(f"No export found. Download the sheet as XLSX into {EXPORT_DIR.relative_to(ROOT)}/ "
                 "or pass --input FILE.xlsx")
    return exports[-1]


def run(path, geocode_on=True, out_path=OUT, quiet=False):
    """Build the site data from one export. Returns (ok, venues)."""
    say = (lambda *a: None) if quiet else print
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    problems = Problems()
    guide = read_field_guide(wb)
    rows = read_venues(wb, guide, problems)
    seen = validate(rows, guide, problems)
    created = assign_ids(rows, seen)
    looked_up = geocode(rows, problems, geocode_on) if not problems.errors else []

    for w in problems.warnings:
        say(f"warning  {w}")
    if problems.errors:
        say(f"\n{len(problems.errors)} problem(s) to fix in the sheet before publishing:")
        for e in problems.errors:
            say(f"  {e}")
        say("\nNothing was written.")
        return False, None

    history, candidates = load_history()
    venues = publish(rows, guide, history)

    old = json.loads(out_path.read_text("utf-8"))["venues"] if out_path.exists() else []
    changes, moved = summarize_changes(old, venues)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps({
        "syncedAt": datetime.date.today().isoformat(),
        "options": {PUBLIC_KEYS[c]: guide[c]["allowed"] for c in FILTER_COLS},
        "venues": venues,
    }, indent=1, ensure_ascii=False) + "\n", "utf-8")

    say(f"Read {len(rows)} rows from {path.name}; published {len(venues)} Approved venues "
        f"to {out_path.relative_to(ROOT) if out_path.is_relative_to(ROOT) else out_path}.")
    if created:
        say("\nNew Venue IDs. Paste these into the sheet so they stay the same:")
        for n, name, vid in created:
            say(f"  Row {n}  {name}  ->  {vid}")
    if looked_up:
        say("\nCoordinates looked up from addresses (add them to the sheet):")
        for n, vid, addr, (la, lo) in looked_up:
            say(f"  Row {n}  {vid}: {la}, {lo}  ({addr})")
    say("\nChanges since the last sync:" if changes else "\nNo changes since the last sync.")
    for line in changes + moved:
        say(line)
    no_history = [v["id"] for v in venues if not v["history"]]
    if no_history:
        say(f"\n{len(no_history)} venue(s) with no Off Tech history yet: {', '.join(no_history)}")
    if candidates:
        say("\nPlaces with several Off Tech listings that aren't on the sheet:")
        for c, name, n in candidates:
            say(f"  {n:>3}  {name}")
    return True, venues


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--input", help="exported .xlsx (default: newest in sheet-export/)")
    ap.add_argument("--no-geocode", action="store_true", help="don't look up missing coordinates")
    a = ap.parse_args()
    ok, _ = run(pick_input(a.input), geocode_on=not a.no_geocode)
    sys.exit(0 if ok else 1)
