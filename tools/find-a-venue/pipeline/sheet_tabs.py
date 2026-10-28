"""Write the newsletter-generated tabs for the collaborator sheet.

    python tools/find-a-venue/pipeline/sheet_tabs.py                 # -> data/off_tech_tabs.xlsx
    python tools/find-a-venue/pipeline/sheet_tabs.py --into FILE.xlsx  # replace the tabs in FILE

Tabs:
  Off Tech history  one row per venue: listings, issues, top activities, host...
  Past events       one row per event held at a venue, newest first, with the
                    issue that listed it. Private and neighborhood-only
                    addresses are never included.

Copy those two tabs into the Proton Sheet to refresh them. This script never
touches the Venues or Field guide tabs, and never talks to the sheet.
"""
import argparse, csv, sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import history  # noqa: E402
from activities import series, tag  # noqa: E402
from data import ISSUES, rows  # noqa: E402

DATA = HERE.parent / "data"
DEFAULT_OUT = DATA / "off_tech_tabs.xlsx"
ARCHIVE = "https://buttondown.com/nycofftech/archive/"

NOTE = ("Generated from the NYC Off Tech newsletter archive by "
        "tools/find-a-venue/pipeline/sheet_tabs.py. Don't edit by hand: "
        "re-import this tab after each new issue.")

PAST_COLUMNS = [("Venue ID", 34), ("Venue name", 36), ("Event date", 12),
                ("Event", 60), ("Host (as listed)", 36), ("Activity", 28),
                ("Listed in issue", 14), ("Issue link", 60)]
HISTORY_WIDTHS = [34, 12, 12, 14, 14, 40, 60, 60, 36]


def venue_names():
    """Venue ID -> name as it appears on the sheet (seed), else the pipeline's name."""
    names = {}
    with open(HERE / "venue_ids.csv", encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            names[r["venue_id"]] = r["pipeline_name"]
    ws = openpyxl.load_workbook(history.SEED, read_only=True)["Venues"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] and r[2]:
            names[r[0]] = r[2]
    return names


def past_events(venue_order):
    """One row per (venue, event, date). A listing repeated in a later issue
    keeps the first issue that listed it."""
    ids = history.load_ids()
    names = venue_names()
    seen, out = set(), []
    for r in sorted(rows(), key=lambda r: ISSUES[r["issue"]][0]):
        vid = ids.get(r["vcode"])
        if not vid or r["vcode"] in history.EXCLUDED:
            continue
        key = (vid, series(r["event"]).lower(), r["date"])
        if key in seen:
            continue
        seen.add(key)
        issue_date, _, slug = ISSUES[r["issue"]]
        out.append([vid, names.get(vid, vid), r["date"], r["event"], r["host"] or "",
                    tag(r["event"]), issue_date, f"{ARCHIVE}{slug}/"])
    # Venues in the history tab's order (most listings first), newest event first.
    order = {vid: i for i, vid in enumerate(venue_order)}
    out.sort(key=lambda x: x[2], reverse=True)
    out.sort(key=lambda x: order.get(x[0], len(order)))
    return out


def write_tab(wb, title, header, rows_, widths, index=None):
    if title in wb.sheetnames:
        index = wb.sheetnames.index(title)
        del wb[title]
    ws = wb.create_sheet(title, index)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="555555")
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows_:
        ws.append(r)
    ws.append([])
    ws.append([NOTE])
    ws.cell(ws.max_row, 1).font = Font(italic=True, color="555555")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    return ws


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--into", help="workbook whose two tabs to replace (default: write a new file)")
    a = ap.parse_args()

    hist, _ = history.build()
    history.write(hist)
    past = past_events([h["Venue ID"] for h in hist])

    if a.into:
        path = Path(a.into)
        wb = openpyxl.load_workbook(path)
    else:
        path = DEFAULT_OUT
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    write_tab(wb, "Off Tech history", history.COLUMNS,
              [[h[c] for c in history.COLUMNS] for h in hist], HISTORY_WIDTHS)
    past_index = wb.sheetnames.index("Off Tech history") + 1 if "Past events" not in wb.sheetnames else None
    write_tab(wb, "Past events", [c for c, _ in PAST_COLUMNS], past,
              [w for _, w in PAST_COLUMNS], past_index)
    wb.save(path)
    print(f"Wrote {len(hist)} venues and {len(past)} past events to {path}")


if __name__ == "__main__":
    main()
