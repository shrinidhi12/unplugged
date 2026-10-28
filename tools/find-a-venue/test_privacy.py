"""Fails if anything the Field guide keeps off the site reaches the site's data.

    npm run venues:test
"""
import json, re, shutil, sys, tempfile, unittest
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import sync  # noqa: E402

SECRET = "SECRET"


def guide():
    return sync.read_field_guide(openpyxl.load_workbook(sync.SEED, read_only=True))


class FieldGuideMapping(unittest.TestCase):
    def test_only_shown_columns_have_public_keys(self):
        g = guide()
        for col in sync.PUBLIC_KEYS:
            self.assertEqual(g[col]["shown"], "Yes", f"{col} is published but not marked 'Shown on site = Yes'")

    def test_contact_keys_are_collaborator_columns(self):
        g = guide()
        for col in sync.CONTACT_KEYS:
            self.assertEqual(g[col]["shown"], "Collaborators only", col)
        self.assertNotIn("Relationship notes", sync.CONTACT_KEYS)


class BuiltData(unittest.TestCase):
    """Run a sync on a copy of the seed with every private cell filled in."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        src = cls.tmp / "export.xlsx"
        shutil.copy(sync.SEED, src)
        wb = openpyxl.load_workbook(src)
        ws = wb["Venues"]
        header = [c.value for c in ws[1]]
        g = sync.read_field_guide(wb)
        private = [c for c, v in g.items() if v["shown"] != "Yes" and c not in ("Venue ID", "Status")]
        cls.ok_row = 2
        for r in range(2, ws.max_row + 1):
            if not ws.cell(r, 1).value:
                continue
            for c in private:
                cell = ws.cell(r, header.index(c) + 1)
                cell.value = "Yes" if (c == sync.OK_CONTACT and r == cls.ok_row) else (
                    "No" if c == sync.OK_CONTACT else f"{SECRET}:{c}:{r}")
            b = ws.cell(r, header.index("Borough") + 1)
            b.value = b.value or "Manhattan"
        wb.save(src)
        cls.out = cls.tmp / "venues.json"
        ok, cls.venues = sync.run(src, geocode_on=False, out_path=cls.out, quiet=True)
        assert ok, "sync failed on the test export"
        cls.text = cls.out.read_text("utf-8")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_no_private_values_except_agreed_contact(self):
        leaks = set(re.findall(rf"{SECRET}:([^:\"]+):(\d+)", self.text))
        allowed = {(c, str(self.ok_row)) for c in sync.CONTACT_KEYS}
        self.assertEqual(leaks - allowed, set(), "private cells reached venues.json")

    def test_agreed_contact_is_published(self):
        with_contact = [v for v in self.venues if "contact" in v]
        self.assertEqual(len(with_contact), 1)
        self.assertEqual(set(with_contact[0]["contact"]), set(sync.CONTACT_KEYS.values()))

    def test_relationship_notes_never_published(self):
        self.assertNotIn("Relationship notes", self.text)


class CommittedData(unittest.TestCase):
    """The venues.json the site actually ships."""

    def test_only_known_keys(self):
        data = json.loads(sync.OUT.read_text("utf-8"))
        allowed = {"id", "history", "contact", *sync.PUBLIC_KEYS.values()}
        for v in data["venues"]:
            self.assertLessEqual(set(v), allowed, v["id"])
            self.assertLessEqual(set(v.get("contact", {})), set(sync.CONTACT_KEYS.values()), v["id"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
